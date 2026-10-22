from pptx import Presentation

prs = Presentation()

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Auntie's PG Website Pitch"
subtitle.text = "A Professional Web Solution for Paying Guest Accommodation Management"

# Overview slide
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Project Overview'
tf = body_shape.text_frame
tf.text = 'Auntie\'s PG is a comprehensive web application designed to showcase and manage a paying guest accommodation facility.'

p = tf.add_paragraph()
p.text = 'Built using Flask framework with SQLite database for data storage.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Features public-facing pages for visitors and a secure admin panel for management.'
p.level = 1

# Features slide
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Key Features'
tf = body_shape.text_frame
tf.text = 'Public Pages:'

p = tf.add_paragraph()
p.text = 'Home: Welcome page with facility overview.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Gallery: Photo gallery of rooms and facilities, categorized by type.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Residents: Showcase current residents with photos and details.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Contact: Form for inquiries, stored securely in database.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Admin Panel: Secure login for management.'
p.level = 2

p = tf.add_paragraph()
p.text = 'Dashboard: View and manage residents, gallery images, and contact inquiries.'
p.level = 2

# Technical Stack
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Technical Stack'
tf = body_shape.text_frame
tf.text = 'Backend: Flask (Python web framework)'

p = tf.add_paragraph()
p.text = 'Database: SQLite for lightweight, file-based storage.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Frontend: HTML, CSS, JavaScript with Jinja2 templating.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Security: CSRF protection, secure file uploads, admin authentication.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Deployment: Ready for Heroku or similar with Procfile and gunicorn.'
p.level = 1

# Benefits
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Benefits'
tf = body_shape.text_frame
tf.text = 'User-Friendly: Intuitive interface for visitors and admins.'

p = tf.add_paragraph()
p.text = 'Secure: Implements best practices for web security.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Scalable: Easy to add new features or expand.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Cost-Effective: Built with open-source tools, low maintenance.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Professional: Clean design suitable for business presentation.'
p.level = 1

# Next Steps
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]

title_shape.text = 'Next Steps'
tf = body_shape.text_frame
tf.text = 'Demo: Live demonstration of the application.'

p = tf.add_paragraph()
p.text = 'Customization: Tailor design and features to client needs.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Deployment: Set up on hosting platform.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Training: Guide on using the admin panel.'
p.level = 1

p = tf.add_paragraph()
p.text = 'Support: Ongoing maintenance and updates.'
p.level = 1

prs.save('aunties_pg_pitch.pptx')

print("PPT created successfully!")